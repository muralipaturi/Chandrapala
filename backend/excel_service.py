from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from models import TestCase


def create_test_case_excel(
    test_cases: list[TestCase]
):

    # =====================================================
    # CREATE WORKBOOK
    # =====================================================

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Test Cases"


    # =====================================================
    # HEADERS
    # =====================================================

    headers = [
        "Test Case ID",
        "Type",
        "Title",
        "Priority",
        "Preconditions",
        "Test Data",
        "Steps",
        "Expected Result",
        "Automation Candidate"
    ]

    worksheet.append(headers)


    # =====================================================
    # HEADER FORMATTING
    # =====================================================

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # =====================================================
    # TEST CASE DATA
    # =====================================================

    for test_case in test_cases:

        worksheet.append([
            test_case.id,

            test_case.type,

            test_case.title,

            test_case.priority,

            "\n".join(
                test_case.preconditions
            ),

            "\n".join(
                test_case.test_data
            ),

            "\n".join(
                test_case.steps
            ),

            test_case.expected_result,

            (
                "Yes"
                if test_case.automation_candidate
                else "No"
            )
        ])


    # =====================================================
    # FORMAT CELLS
    # =====================================================

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


    # =====================================================
    # COLUMN WIDTHS
    # =====================================================

    column_widths = {

        1: 16,

        2: 18,

        3: 45,

        4: 12,

        5: 40,

        6: 40,

        7: 55,

        8: 55,

        9: 22
    }


    for column_number, width in column_widths.items():

        column_letter = (
            get_column_letter(
                column_number
            )
        )

        worksheet.column_dimensions[
            column_letter
        ].width = width


    # =====================================================
    # FREEZE HEADER
    # =====================================================

    worksheet.freeze_panes = "A2"


    # =====================================================
    # FILTER
    # =====================================================

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )


    return workbook